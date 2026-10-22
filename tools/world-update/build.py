"""수집한 공개 통계로 abyss/js/world-data.js 의 해당 항목을 새로 덮어쓴다.

사용: python3 -I build.py <수집 폴더> <world-data.js>
  수집 폴더 구조(수집 방법은 같은 폴더의 README.md):
    fao/{Production_Crops_Livestock,Trade_CropsLivestock}_E_All_Data_(Normalized).zip
    un/WPP2024_*.csv.gz · un/wup/WUP2025-F15-*.xlsx · un/undesa_pd_2024_ims_stock_by_sex_and_destination.xlsx
    wb/*.json · unhcr/unhcr_2024.json · owid/owid-energy-data.csv
    other/ember_yearly_full_release_long_format.csv · irena/*.csv · wmd/6.4_by_minerals.xlsx
    iea/iea.xlsx (IEA World Energy Balances Highlights 2026판) · ei/ei.xlsx (Energy Institute Statistical Review 2025판 전체 데이터) · wto/WtoData.csv · unsd/Download-GDPcurrent-USD-countries.xlsx
덮어쓰는 항목은 아래 각 단계에 연도와 출처를 적어 두었다. 0·결측은 키를 만들지 않는다(world-data.js 머리말 규칙).
"""
import csv,gzip,io,json,os,re,sys,zipfile
import openpyxl,pycountry
SRC,PATH=sys.argv[1],sys.argv[2]
P=lambda *a:os.path.join(SRC,*a)
txt=open(PATH,encoding='utf-8').read()
m=re.search(r'^const WORLD_DATA=(\{.*\});\s*$',txt,flags=re.M)
W=json.loads(m.group(1))
EXTRA={'XKX':'xk'}
def iso2(a3):
    if a3 in EXTRA:return EXTRA[a3]
    c=pycountry.countries.get(alpha_3=a3)
    return c.alpha_2.lower() if c else None
def m49_iso2(x):
    c=pycountry.countries.get(numeric=str(x).strip("' ").zfill(3))
    return c.alpha_2.lower() if c else None
def put(k,i,v,nd=2):
    if i not in W or v is None:return False
    v=round(float(v),nd)
    if v==0:W[i].pop(k,None);return False
    W[i][k]=v;return True
OLD={i:dict(v) for i,v in W.items()}
log=[]
def done(k,n,note=''):log.append(f'{k:5s} {n:4d}개국 {note}')

# ── UN WPP 2024 (Medium) : 연령 구조 2025 · 합계출산율 2025 · 출생/사망률 2023 · 순이동률 2025 ──
for k in ('y0','y1','y2','tfr'):
    for i in W:W[i].pop(k,None)
age={}
with gzip.open(P('un','WPP2024_PopulationByAge5GroupSex_Medium.csv.gz'),'rt',encoding='utf-8-sig') as h:
    for r in csv.DictReader(h):
        if r['Time']!='2025' or r['LocTypeName']!='Country/Area' or not r['ISO3_code']:continue
        i=iso2(r['ISO3_code']);
        if not i:continue
        s=int(r['AgeGrpStart']);a=age.setdefault(i,[0,0,0])
        a[0 if s<15 else 1 if s<65 else 2]+=float(r['PopTotal'])
n=0
for i,a in age.items():
    t=sum(a)
    if t and i in W:
        for j,k in enumerate(('y0','y1','y2')):put(k,i,a[j]/t*100)
        n+=1
done('y0-2',n,'UN WPP 2024, 2025년')
for k in('cbr','cdr','nmr'):
    for i in W:W[i].pop(k,None)
cnt=dict(tfr=0,cbr=0,cdr=0,nmr=0)
with gzip.open(P('un','WPP2024_Demographic_Indicators_Medium.csv.gz'),'rt',encoding='utf-8-sig') as h:
    for r in csv.DictReader(h):
        if r['LocTypeName']!='Country/Area' or not r['ISO3_code']:continue
        i=iso2(r['ISO3_code'])
        if not i:continue
        if r['Time']=='2025':
            if r['TFR'] and put('tfr',i,r['TFR']):cnt['tfr']+=1
            if r['CNMR'] and put('nmr',i,r['CNMR'],1):cnt['nmr']+=1
        if r['Time']=='2023':
            if r['CBR'] and put('cbr',i,r['CBR'],1):cnt['cbr']+=1
            if r['CDR'] and put('cdr',i,r['CDR'],1):cnt['cdr']+=1
for k,v in cnt.items():done(k,v,'UN WPP 2024 '+('2023년' if k in('cbr','cdr') else '2025년'))

# ── UN WUP 2025 : 도시화율(국가별 도시 정의, 2025년) ──
ws=openpyxl.load_workbook(P('un','wup','WUP2025-F15-National-Definitions_percPop_by_category.xlsx'),read_only=True)['Urban']
rows=list(ws.iter_rows(values_only=True))
hi=next(i for i,r in enumerate(rows) if r and r[0]=='Index');hdr=rows[hi]
yc=next(j for j,x in enumerate(hdr) if str(x)=='2025')
for i in W:W[i].pop('ur',None)
n=0
for r in rows[hi+1:]:
    if r[4] and put('ur',iso2(r[4]) if iso2(r[4]) else '',r[yc]):n+=1
done('ur',n,'UN WUP 2025, 2025년(국가별 도시 정의)')

# ── World Bank : 산업 구조 · 연료/제조업 비중 (2022~2025 중 나라별 가장 최근) ──
def wb_latest(code):
    out={}
    for r in json.load(open(P('wb',code+'.json'),encoding='utf-8')):
        a3=r.get('countryiso3code')
        if not a3 or r['value'] is None:continue
        i=iso2(a3)
        if not i:continue
        y=int(r['date'])
        if i not in out or y>out[i][0]:out[i]=(y,r['value'])
    return out
for k,code in (('i1','NV.AGR.TOTL.ZS'),('i2','NV.IND.TOTL.ZS'),('i3','NV.SRV.TOTL.ZS'),('fuel','TX.VAL.FUEL.ZS.UN'),('manf','NV.IND.MANF.ZS')):
    for i in W:W[i].pop(k,None)
    d=wb_latest(code);n=sum(put(k,i,v) for i,(y,v) in d.items());done(k,n,f'World Bank {code}, 나라별 최근({min(y for y,_ in d.values())}~{max(y for y,_ in d.values())})')

# ── Energy Institute Statistical Review 2025판(자료 2024년) : 생산 · 소비 · 무역 ──
EI=openpyxl.load_workbook(P('ei','ei.xlsx'),read_only=True,data_only=True)
EI_NAME={'US':'us','Russia':'ru','Russian Federation':'ru','Republic of Congo':'cg','Turkey':'tr','Iran':'ir','South Korea':'kr','Taiwan':'tw','Vietnam':'vn','Trinidad & Tobago':'tt',
         'China Hong Kong SAR':'hk','Czech Republic':'cz','North Macedonia':'mk','Brunei':'bn','Congo':'cg','Democratic Republic of Congo':'cd',
         'Ivory Coast':'ci',"Cote d'Ivoire":'ci','Laos':'la','Syria':'sy','Bolivia':'bo','Venezuela':'ve','Tanzania':'tz','United Kingdom':'gb',
         'Slovak Republic':'sk','Kyrgyzstan':'kg','Republic of Korea':'kr'}
def ei_iso(nm):
    nm=re.sub(r'\s+',' ',nm.strip().rstrip('*^†').strip())
    if nm in('USSR','EU','West Africa','East & S. Africa','Republic of Congo' and 'zzz'):return None
    if nm in EI_NAME:return EI_NAME[nm]
    if re.match(r'^(Other|Total|Rest|Middle|Non-|OECD|European|CIS|Asia|Africa|Europe|North|South & |S\. |Central|Eastern|Western|Australasia|Persian)',nm):return None
    try:return pycountry.countries.lookup(nm).alpha_2.lower()
    except LookupError:
        print('EI 이름 매칭 실패(건너뜀):',nm);return None
def ei_table(name):
    rows=list(EI[name].iter_rows(values_only=True))
    hi=next(i for i,r in enumerate(rows) if r and len(r)>1 and isinstance(r[1],(int,float)) and r[1]>=1960)
    yc=[j for j,x in enumerate(rows[hi]) if x==2024][0]
    out={}
    for r in rows[hi+1:]:
        if not r[0]:continue
        k=str(r[0]).strip()
        if k.startswith('Total World'):break
        if isinstance(r[yc],(int,float)):out[k]=r[yc]
    return out
for k in('cp','gp','op','cc','gc','oc'):
    for i in W:W[i].pop(k,None)
EJ=277.78                    # 1 EJ = 277.78 TWh
EIS=[('cp','Coal Production - EJ',EJ),('gp','Gas Production - EJ',EJ),('op','Oil Production - tonnes',11.63),   # 석유: 1 Mt = 11.63 TWh
     ('cc','Coal Consumption - EJ',EJ),('gc','Gas Consumption - EJ',EJ),('oc','Oil Consumption - EJ',EJ)]
cn={}
for k,sheet,mul in EIS:
    n=0
    for nm,v in ei_table(sheet).items():
        i=ei_iso(nm)
        if i and put(k,i,v*mul,1):n+=1
    cn[k]=n
kept={k:0 for k in cn}
for k in cn:
    for i,v in OLD.items():
        if i in W and k not in W[i] and k in v:W[i][k]=v[k];kept[k]+=1   # EI에 따로 없는 나라는 전에 받은 값(EIA·OWID)을 그대로 둔다
for k,v in cn.items():done(k,v,f'Energy Institute 2025판 2024년 (+이전 값 유지 {kept[k]}개국)')

# ── Energy Institute 무역 2024년 : 석탄 · 천연가스 · 원유 (EI가 나라별로 따로 주는 곳만 덮어씀) ──
for k in('cx','cm','gx','gm'):
    pass
rows=list(EI['Coal - Trade movements'].iter_rows(values_only=True))
hi=next(i for i,r in enumerate(rows) if r and r[0]=='Exajoules');yc=[j for j,x in enumerate(rows[hi]) if x==2024][0];sec=None;ct={'cx':0,'cm':0};EI_COAL={'cx':set(),'cm':set()}
for r in rows[hi+1:]:
    if not r[0]:continue
    t=str(r[0]).strip()
    if t in('Imports','Exports'):sec=t;continue
    if t.startswith('Total World'):continue
    if sec and isinstance(r[yc],(int,float)):
        i=ei_iso(t)
        if i and i in W:
            k='cm' if sec=='Imports' else 'cx'
            W[i].pop(k,None)
            if put(k,i,r[yc]*EJ,1):ct[k]+=1;EI_COAL[k].add(i)
done('cx',ct['cx'],'Energy Institute 석탄 무역 2024(EI가 나라별로 주는 나라만)');done('cm',ct['cm'],'Energy Institute 석탄 무역 2024')
# IEA World Energy Balances Highlights 2026판 2024년 : 석탄 수출입 — EI가 나라별로 주지 않는 나라만 채운다(PJ → TWh = ÷3.6)
IEA_NAME={"People's Republic of China":'cn','Korea':'kr','Republic of Turkiye':'tr','Viet Nam':'vn','Czech Republic':'cz','Slovak Republic':'sk'}
IEA_SKIP={'OECD Total','Non-OECD Total','Africa','Non-OECD Americas','Non-OECD Asia (including China)','Non-OECD Europe and Eurasia','Middle East','IEA Total','IEA and Accession/Association countries','World'}
ie=openpyxl.load_workbook(P('iea','iea.xlsx'),read_only=True,data_only=True)['TimeSeries_1971-2025']
irows=list(ie.iter_rows(values_only=True));yi=[j for j,h in enumerate(irows[1]) if str(h)=='2024'][0]
ci={'cx':0,'cm':0}
for r in irows[2:]:
    if not r[0] or r[0] in IEA_SKIP or r[1]!='Coal, peat and oil shale' or r[2] not in('Imports (PJ)','Exports (PJ)'):continue
    i=IEA_NAME.get(r[0])
    if i is None:
        try:i=pycountry.countries.lookup(r[0]).alpha_2.lower()
        except LookupError:print('IEA 이름 매칭 실패(건너뜀):',r[0]);continue
    v=r[yi]
    if not isinstance(v,(int,float)) or i not in W:continue
    k='cm' if r[2].startswith('Imports') else 'cx'
    if i in EI_COAL.get(k,set()):continue          # EI가 나라별로 준 나라는 EI 값을 쓴다
    W[i].pop(k,None)
    if put(k,i,abs(v)/3.6,1):ci[k]+=1
done('cx',ci['cx'],'IEA World Energy Balances Highlights 2026판 2024년 (EI에 없는 나라)');done('cm',ci['cm'],'IEA 2026판 2024년 (EI에 없는 나라)')
def gas_matrix(name):
    rows=list(EI[name].iter_rows(values_only=True))
    hi=next(i for i,r in enumerate(rows) if r and r[0] and str(r[0]).strip().startswith(('To','Billion cubic metres To')) and 'Total imports' in [str(x).strip() for x in r if x])
    names=[]
    for j in range(len(rows[hi])):
        top=rows[hi-1][j] if hi>0 else None
        cell=rows[hi][j]
        names.append((str(top or '')+' '+str(cell or '')).strip() if (top and cell and j>0 and str(rows[hi-1][j]).strip() not in('Billion cubic metres',)) else str(cell or top or '').strip())
    exp={};imp={}
    ti=[j for j,x in enumerate(rows[hi]) if x and str(x).strip()=='Total imports'][0]
    for r in rows[hi+1:]:
        if not r[0]:continue
        t=str(r[0]).strip()
        if t=='Total exports':
            for j,x in enumerate(r):
                if j>0 and j<ti and isinstance(x,(int,float)):exp[names[j]]=x
            break
        if isinstance(r[ti],(int,float)):imp[t]=r[ti]
    return exp,imp
GX={};GM={}
for nm in('Gas trade 2024 - LNG','Gas trade 2024 - pipeline'):
    ex,im=gas_matrix(nm)
    for k,v in ex.items():GX[k]=GX.get(k,0)+v
    for k,v in im.items():GM[k]=GM.get(k,0)+v
# 나라별 합계는 '지역 간' 행렬이라 같은 지역 안 거래(러시아→다른 CIS 등)가 빠진다 — EI가 'Gas - Trade movements'에서
# 따로 합계를 주는 나라(미국·브라질·러시아·중국·인도)는 그 표의 값을 쓴다.
rows=list(EI['Gas - Trade movements'].iter_rows(values_only=True))
yg=[j for j,x in enumerate(rows[2]) if x==2024][0];cur=None;GT={}
for r in rows[3:]:
    if not r[0]:continue
    t=str(r[0]).strip()
    if all(x is None for x in r[1:6]):cur=t;continue
    if t in('Total imports','Total exports') and isinstance(r[yg],(int,float)):GT[(cur,t)]=r[yg]
for (c,t),v in GT.items():
    if c in('US','Brazil','Russian Federation','China','India'):
        (GM if t=='Total imports' else GX)[c]=v
# 천연가스 수출입은 EI가 나라별로 주는 나라만 둔다(수능특강 순위와 같은 기준) — 독일 같은 'Other EU'는 값이 없다
for k in('gx','gm'):
    for i in W:W[i].pop(k,None)
cg={'gx':0,'gm':0}
for k,src in(('gx',GX),('gm',GM)):
    for nm,v in src.items():
        nm2=re.sub(r'^(From|To)\s+','',nm.replace('*',''))
        i=ei_iso(nm2)
        if i and i in W and v>0 and put(k,i,v*10,1):cg[k]+=1     # 1 bcm = 10 TWh (EI 생산 환산과 같다)
done('gx',cg['gx'],'Energy Institute 가스 무역 2024(LNG+파이프라인, 나라별로 주는 나라만)');done('gm',cg['gm'],'Energy Institute 가스 무역 2024')

# 원유 수출입(TWh, 2024): CEPII BACI(원유 270900) 바탕 + Energy Institute 2025판이 나라별로 주는 값을 덮어쓴다 (1 Mt = 11.63 TWh)
bc=json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)),'baci2024.json')))
for i in W:W[i].pop('ox',None);W[i].pop('om',None)
oil={'ox':{},'om':{}}
for key,val in bc.items():
    f,a3=key.split('|');i=iso2(a3)
    if i and i in W and val>0:oil['ox' if f=='X' else 'om'][i]=val
rows=list(EI['Oil - Trade movements in 23-24'].iter_rows(values_only=True))
for r in rows[4:]:
    if not r[0]:continue
    t=str(r[0]).strip()
    if t=='Thousand barrels daily':break
    i=ei_iso(t)
    if i and i in W and isinstance(r[5],(int,float)) and isinstance(r[7],(int,float)):
        if r[7]>0:oil['ox'][i]=r[7]        # 2024 원유 수출(백만 톤)
        if r[5]>0:oil['om'][i]=r[5]        # 2024 원유 수입
oil['ox']['ir']=78.0       # 이란: OPEC ASB 2025 원유 수출 1,566천 b/d ≈ 78 Mt (BACI에 안 잡힌다)
for k,d in oil.items():
    n=0
    for i,v in d.items():
        if put(k,i,v*11.63,1):n+=1
    done(k,n,'CEPII BACI + Energy Institute 2025판 원유 무역 2024')

# ── 발전량(TWh, 2024년) : Ember(석탄·가스·원자력·수력·태양광·풍력·바이오) + IRENA(석유·지열·해양) ──
# 순서: [석탄,가스,석유,원자력,수력,태양광,풍력,바이오,지열,해양]
EMB={'Coal':0,'Gas':1,'Nuclear':3,'Hydro':4,'Solar':5,'Wind':6,'Bioenergy':7}
el={}
for r in csv.DictReader(open(P('other','ember_yearly_full_release_long_format.csv'),encoding='utf-8-sig')):
    if r['Year']!='2024' or r['Category']!='Electricity generation' or r['Unit']!='TWh' or r['Subcategory']!='Fuel':continue
    if r['Variable'] not in EMB or not r['ISO 3 code'] or r['Area type']!='Country or economy':continue
    i=iso2(r['ISO 3 code'])
    if not i or i not in W:continue
    el.setdefault(i,[0.0]*10)[EMB[r['Variable']]]=float(r['Value'] or 0)
IR={'Oil':2,'Geothermal energy':8,'Marine energy':9}
irf=[f for f in os.listdir(P('irena')) if f.endswith('.csv')][0]
IRC=json.load(open(P('irena','country_codes.json'),encoding='utf-8'))   # 표시 이름 → ISO3 (IRENA PxWeb 메타데이터)
for r in csv.DictReader(open(P('irena',irf),encoding='latin-1')):
    if r['Year']!='2024' or r['Technology'] not in IR or not r['Data Type'].startswith('Electricity Generation'):continue
    v=r['Electricity statistics']
    if v in('..',''):continue
    a3=IRC.get(r['Country/area'])
    i=iso2(a3) if a3 else None
    if not i:continue
    if i in W and i in el:el[i][IR[r['Technology']]]=float(v)/1000   # GWh → TWh
n=0
kept=0
for i in W:
    W[i].pop('el',None);W[i].pop('el_year',None)
    if i in el and sum(el[i])>0:W[i]['el']=[round(x,2) for x in el[i]];W[i]['el_year']=2024;n+=1
    elif 'el' in OLD[i]:W[i]['el']=OLD[i]['el'];W[i]['el_year']=OLD[i].get('el_year');kept+=1   # Ember에 없는 나라는 IRENA 2023 유지
done('el',n,f'Ember 2024 + IRENA 2024(석유·지열·해양) (+IRENA 2023 유지 {kept}개국)')

# ── FAOSTAT 2024 : 생산 · 사육두수 · 곡물 수출입 (+ 새 작물·육류) ──
def fao_rows(zname,items,elements,years=('2024',)):
    z=zipfile.ZipFile(P('fao',zname));nm=next(x for x in z.namelist() if x.endswith('(Normalized).csv'))
    with io.TextIOWrapper(z.open(nm),encoding='latin-1') as h:
        for r in csv.DictReader(h):
            if r['Year'] in years and r['Item Code'] in items and r['Element'] in elements and r['Value']:
                yield r
PRODK={'15':'wh','27':'ri','56':'co','656':'cof','328':'ctn','125':'cas','661':'cca','667':'tea','836':'rub','257':'pal','156':'suc','260':'oli','560':'grp','577':'dat','867':'bef','1035':'prk','977':'mtn'}
HEAD={'866':'ct','976':'sh','1034':'pg'}
allk=list(PRODK.values())+list(HEAD.values())
for k in allk:
    for i in W:W[i].pop(k,None)
cn={k:0 for k in allk}
for r in fao_rows('Production_Crops_Livestock_E_All_Data_(Normalized).zip',set(PRODK)|set(HEAD),{'Production','Stocks'},years=('2024','2023')):
    i=m49_iso2(r['Area Code (M49)'])
    if not i:continue
    ic=r['Item Code'];v=float(r['Value'])
    if r['Year']=='2023' and ic!='257':continue   # 팜유만 2023년(FAOSTAT에 2024년 값이 아직 없다)
    if ic in HEAD and r['Element']=='Stocks' and r['Unit']=='An':k=HEAD[ic]
    elif ic in PRODK and r['Element']=='Production' and r['Unit']=='t':k=PRODK[ic]
    else:continue
    if put(k,i,v/1e4):cn[k]+=1
for k,v in cn.items():done(k,v,'FAOSTAT 2024(만 톤·만 두)')
TR={'15':('whx','whm'),'30':('rix','rim'),'56':('cox','com')}   # 쌀 무역은 '도정 환산 쌀'(30)
for k in [x for p in TR.values() for x in p]:
    for i in W:W[i].pop(k,None)
cn={x:0 for p in TR.values() for x in p}
for r in fao_rows('Trade_CropsLivestock_E_All_Data_(Normalized).zip',set(TR),{'Export quantity','Import quantity'}):
    i=m49_iso2(r['Area Code (M49)'])
    if not i or r['Unit']!='t':continue
    k=TR[r['Item Code']][0 if r['Element'].startswith('Export') else 1]
    if put(k,i,float(r['Value'])/1e4):cn[k]+=1
for k,v in cn.items():done(k,v,'FAOSTAT 무역 2024(만 톤)')

# ── World Mining Data 2023 : 광물 (다이아몬드는 Kimberley Process 유지) ──
WMD_NAME={'Bosnia-Herzegovina':'ba','Congo, Rep.':'cg','Congo, D.R.':'cd','Korea, North':'kp','Korea, South':'kr','Russia':'ru',"Cote d'Ivoire":'ci'}
MIN={'Iron (Fe)':('ior',1e4),'Gold':('gld',1e3),'Silver':('slv',1e3),'Copper':('cop',1e4),'Cobalt':('cbt',1),
     'Chromium (Cr2O3)':('chr',1e4),'Manganese':('mng',1e4),'Bauxite':('bux',1e4),'Tin':('tin',1)}
wb=openpyxl.load_workbook(P('wmd','6.4_by_minerals.xlsx'),read_only=True,data_only=True)
for nm,(k,div) in MIN.items():
    for i in W:W[i].pop(k,None)
    rows=list(wb[nm].iter_rows(values_only=True));hd=next(r for r in rows if r and r[0]=='Country');yc=list(hd).index('2023')
    n=0
    for r in rows[rows.index(hd)+1:]:
        if not r[0] or r[yc] in(None,''):continue
        nm0=str(r[0]).strip()
        i=WMD_NAME.get(nm0)
        if i is None:
            try:i=pycountry.countries.lookup(nm0).alpha_2.lower()
            except LookupError:
                if nm0!='Total':print('WMD 이름 매칭 실패(건너뜀):',nm0)
                continue
        if put(k,i,float(r[yc])/div):n+=1
    done(k,n,f'World Mining Data 2023 · {nm}')

# ── UNHCR 2024 : 난민 발생(출신국) · 수용국 ──
for k in('rfo','rfh'):
    for i in W:W[i].pop(k,None)
ro={};rh={}
for r in json.load(open(P('unhcr','unhcr_2024.json'),encoding='utf-8')):
    n_=int(r.get('refugees') or 0)
    for key,bag in(('coo_iso',ro),('coa_iso',rh)):
        a3=r.get(key)
        if not a3:continue
        i=iso2(a3)
        if i:bag[i]=bag.get(i,0)+n_
done('rfo',sum(put('rfo',i,v,0) for i,v in ro.items()),'UNHCR 2024(출신국 기준 난민)')
done('rfh',sum(put('rfh',i,v,0) for i,v in rh.items()),'UNHCR 2024(수용국 기준 난민)')

# ── UN 국제 이주자 2024 : 거주국 기준 이주자 수(명) ──
for i in W:W[i].pop('mig',None)
ws=openpyxl.load_workbook(P('un','undesa_pd_2024_ims_stock_by_sex_and_destination.xlsx'),read_only=True,data_only=True)['Table 1']
rows=list(ws.iter_rows(values_only=True));hi=next(i for i,r in enumerate(rows) if r and r[1] and str(r[1]).startswith('Region, development group'))
yc=list(rows[hi]).index(2024);n=0
for r in rows[hi+1:]:
    code=r[4]
    if isinstance(code,int) and code<900:
        i=m49_iso2(code)
        if i and r[yc] not in(None,'') and put('mig',i,float(r[yc]),0):n+=1
done('mig',n,'UN DESA 국제 이주자 2024')

# ── WTO 2024 : 상품 수출입액(억 달러)과 농산물·연료/광산물 비중 ──
for k in('tx','tm','txa','txf','tma','tmf'):
    for i in W:W[i].pop(k,None)
tot={};grp={}
for r in csv.DictReader(open(P('wto','WtoData.csv'),encoding='utf-8-sig',errors='replace')):
    a3=r['Reporting Economy ISO3A Code']
    if not a3 or r['Year']!='2024':continue
    i=iso2(a3)
    if not i or i not in W:continue
    flow='x' if r['Indicator Code']=='ITS_MTV_AX' else 'm'
    try:v=float(r['Value'])
    except ValueError:continue
    if r['Product/Sector Code']=='TO':tot[(flow,i)]=v
    elif r['Product/Sector Code'] in('AG','MI'):grp[(flow,i,r['Product/Sector Code'])]=v
cw={k:0 for k in('tx','tm','txa','txf','tma','tmf')}
for (flow,i),v in tot.items():
    if put('t'+flow,i,v/100,1):cw['t'+flow]+=1     # 백만 달러 → 억 달러
    for code,suf in(('AG','a'),('MI','f')):
        g=grp.get((flow,i,code))
        if g is not None and v>0 and put('t'+flow+suf,i,g/v*100,1):cw['t'+flow+suf]+=1
for k,v in cw.items():done(k,v,'WTO 2024')

# ── UN 통계국 국민계정 2024 : 제조업 부가가치(억 달러, 현재 가격) ──
for i in W:W[i].pop('mva',None)
ws=openpyxl.load_workbook(P('unsd','Download-GDPcurrent-USD-countries.xlsx'),read_only=True,data_only=True).worksheets[0]
rows=list(ws.iter_rows(values_only=True));hi=next(i for i,r in enumerate(rows) if r and r[0]=='CountryID');yc=list(rows[hi]).index(2024);n=0
for r in rows[hi+1:]:
    if r[2]=='Manufacturing (ISIC D)' and isinstance(r[0],int):
        i=m49_iso2(r[0])
        if i and r[yc] and put('mva',i,float(r[yc])/1e8,1):n+=1
done('mva',n,'UN 통계국 국민계정 2024(제조업, ISIC D)')

out=txt[:m.start(1)]+json.dumps(W,ensure_ascii=False,separators=(',',':'))+txt[m.end(1):]
open(PATH,'w',encoding='utf-8').write(out)
print('\n'.join(log))
