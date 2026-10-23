"""도시 통계 — UN 세계 도시화 전망 2025(WUP 2025)의 도시(DEGURBA) 파일로 abyss/js/city-stats.js 를 만든다.

  F18  해마다(1975~2050, 5년 간격) 인구 상위 100개 도시 — 순위·좌표·인구
  F19  2025년 상위 100개 도시의 1975~2050년 인구
  F34  인구 규모 계급별 도시 수(5만~10만 … 1000만 이상) — 세계·대륙
  F35  인구 규모 계급별 도시 인구 — 세계·대륙

  '도시'는 위성 격자로 본 고밀도 시가지(1km²당 1,500명 이상이 이어진 5만 명 이상 지역)다.
  행정 구역 경계와 달라 나라 통계청 값과 다를 수 있다.

  python3 tools/world-update/cities.py <내려받은 폴더> abyss/js/city-stats.js
"""
import json,os,sys
import openpyxl
sys.path.insert(0,os.path.dirname(__file__))
from city_ko import KO
DL,OUT=sys.argv[1],sys.argv[2]
P=lambda f:os.path.join(DL,'un','wup',f)
YEARS=list(range(1975,2051,5))
CONT={900:'world',903:'af',935:'as',908:'eu',904:'la',905:'an',909:'oc'}
D={'years':YEARS,'cities':{},'top':{},'classes':{}}

def city(r):
    code=r[8];c=D['cities'].get(code)
    if not c:
        en=r[9];ko=KO.get(code)
        if not ko and '(' in en:ko=en[en.index('(')+1:en.rindex(')')]
        c=D['cities'][code]={'n':ko or en,'en':en,'iso':(r[7] or '').lower(),'cap':1 if r[10] else 0,
            'll':[round(r[11],3),round(r[12],3)],'pop':{}}
    if r[1]==2025:c['area']=r[15];c['built']=round(r[16]) if r[16] else None
    c['pop'][str(r[1])]=round(r[13])        # 천 명
    return code

for f,keep in (('WUP2025-F18-DEGURBA-100_Largest_Cities.xlsx',True),('WUP2025-F19-DEGURBA-100_Largest_Cities_in_2025_by_time.xlsx',False)):
    ws=openpyxl.load_workbook(P(f),read_only=True,data_only=True)['Data']
    for r in ws.iter_rows(values_only=True,min_row=2):
        if not r[0]:continue
        code=city(r)
        if keep:D['top'].setdefault(str(r[1]),[]).append((r[2],code))
D['top']={y:[c for _,c in sorted(v)] for y,v in D['top'].items()}

CLS=['50k-100k','100k-250k','250k-500k','500k-1 million','1-2.5 millions','2.5-5 millions','5-10 millions','>10 millions']
D['classLabels']=['5만~10만','10만~25만','25만~50만','50만~100만','100만~250만','250만~500만','500만~1000만','1000만 이상']
for f,key,div in (('WUP2025-F34-DEGURBA_Cities_by_City_Size_Class.xlsx','n',1),('WUP2025-F35-DEGURBA_Pop_by_City_Size_Class.xlsx','pop',1e6)):
    wb=openpyxl.load_workbook(P(f),read_only=True,data_only=True)
    for i,sn in enumerate(CLS):
        rows=list(wb[sn].iter_rows(values_only=True));hd=next(j for j,r in enumerate(rows) if r and r[0]=='Index');h=rows[hd]
        yc={y:list(h).index(str(y)) for y in (1975,2000,2025,2050)}
        for r in rows[hd+1:]:
            if r and r[3] in CONT:
                for y,j in yc.items():
                    a=D['classes'].setdefault(CONT[r[3]],{}).setdefault(str(y),{'n':[0]*8,'pop':[0]*8})
                    a[key][i]=round((r[j] or 0)/div,2) if div>1 else int(r[j] or 0)   # 인구는 백만 명
D['source']='UN DESA, World Urbanization Prospects 2025 (DEGURBA 도시)'
open(OUT,'w',encoding='utf-8').write('/* 도시 통계 — tools/world-update/cities.py 가 만든다(직접 고치지 말 것). 인구는 천 명, 계급별 인구는 백만 명 */\nconst CITY_STATS='
    +json.dumps(D,ensure_ascii=False,separators=(',',':'))+';\n')
print('ok',len(D['cities']),'도시',os.path.getsize(OUT),'bytes')
print({k:v['2025']['n'][-1] for k,v in D['classes'].items()})
