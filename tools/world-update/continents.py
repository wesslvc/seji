"""대륙별 통계 — 국가 값을 평균 내지 않고, 원천이 대륙 단위로 직접 내놓은 합계를 쓴다.

사용: python3 -I continents.py <수집 폴더> abyss/js/continent-data.js
  대륙 구분은 UN 지리 지역(아프리카·아시아·유럽·라틴 아메리카·앵글로아메리카=북부 아메리카·오세아니아)이다.
  · 인구·출생/사망·연령·순이동: UN WPP 2024 지역 행        · 도시화: UN WUP 2025 지역 행
  · 국제 이주자: UN DESA 2024 지역 행(거주 지역 · 출신 지역)  · 농축산·곡물 무역: FAOSTAT 지역 행(2023)
  · 전력(Ember)·난민(UNHCR)·광물(World Mining Data): 원천에 대륙 행이 없거나 UN 구분과 달라, 나라 값을 UN 지역에
    넣어 '합산'했다(평균이 아니다). · 에너지: Energy Institute 권역(북아메리카·중남미·유럽·CIS·중동·아프리카·아시아태평양)을 쓰되,
    러시아는 유럽에 넣고, 나머지 CIS와 중동은 '건조 아시아(중동·CIS)' 하나로 합쳤다 — 대륙과 경계가 달라 대륙에 억지로 끼워 맞추지는 않았다.
  '라틴 아메리카'의 FAOSTAT 값은 중앙아메리카+남아메리카+카리브 지역 행의 합이다.
"""
import csv,gzip,io,json,os,re,sys,zipfile
import openpyxl,pycountry
SRC,OUT=sys.argv[1],sys.argv[2]
P=lambda *a:os.path.join(SRC,*a)
C=[('af','아프리카','Africa',903),('as','아시아','Asia',935),('eu','유럽','Europe',908),
   ('la','라틴 아메리카','Latin America and the Caribbean',904),('an','앵글로아메리카','Northern America',905),('oc','오세아니아','Oceania',909)]
ID2K={c[3]:c[0] for c in C}
YRS=[1950,1970,1990,2010,2025]
R=lambda v,n=1:None if v in(None,'') else round(float(v),n)
D={'conts':[{'id':k,'name':n,'en':e} for k,n,e,_ in C],'years':YRS}
def per(f):return {k:f(k) for k,*_ in C}

# ── UN WPP ──
tree={}      # LocID → (ParentID, iso3)
dem={}       # (LocID,year) → row
with gzip.open(P('un','WPP2024_Demographic_Indicators_Medium.csv.gz'),'rt',encoding='utf-8-sig') as h:
    for r in csv.DictReader(h):
        if r['LocTypeName'] in('Country/Area','Geographic region','Subregion','World') and r['LocID'] and r['Time']=='2025':
            tree[int(r['LocID'])]=(int(r['ParentID'] or 0),r['ISO3_code'])
        if int(r['LocID']) in ID2K or r['LocID']=='900':
            dem[(int(r['LocID']),int(r['Time']))]=r
def cont_of(loc):
    seen=0
    while seen<6:
        if loc in ID2K:return ID2K[loc]
        if loc==918:return 'an'      # WPP가 북아메리카 국가들의 상위로 쓰는 '북부 아메리카' 하위 지역 코드(행은 없다)
        if loc not in tree:return None
        loc=tree[loc][0];seen+=1
    return None
iso2cont={iso:cont_of(l) for l,(p,iso) in tree.items() if iso}
W=lambda y,col:float(dem[(900,y)][col])
def series(col,years,div=1,nd=1):
    return {k:{str(y):R(float(dem[(l,y)][col])/div,nd) for y in years} for l,k in ID2K.items()}
pop=series('TPopulation1July',YRS,1000,0)
D['pop']={k:{'series':pop[k],'share2025':R(float(dem[(l,2025)]['TPopulation1July'])/W(2025,'TPopulation1July')*100),
  'density':R(dem[(l,2025)]['PopDensity']),'median':R(dem[(l,2025)]['MedianAgePop']),'tfr':R(dem[(l,2025)]['TFR'],2),'lex':R(dem[(l,2025)]['LEx']),
  'cbr':{str(y):R(dem[(l,y)]['CBR']) for y in(1950,1970,1990,2010,2023)},'cdr':{str(y):R(dem[(l,y)]['CDR']) for y in(1950,1970,1990,2010,2023)},
  'nat':{str(y):R(dem[(l,y)]['NatChangeRT']) for y in(1950,1970,1990,2010,2023)},
  'netmig':{str(y):R(dem[(l,y)]['NetMigrations'],0) for y in YRS},'cnmr':R(dem[(l,2025)]['CNMR']),
  'growth':R(dem[(l,2025)]['PopGrowthRate'],2)} for l,k in ID2K.items()}
age={}
with gzip.open(P('un','WPP2024_PopulationByAge5GroupSex_Medium.csv.gz'),'rt',encoding='utf-8-sig') as h:
    for r in csv.DictReader(h):
        if r['Time']=='2025' and int(r['LocID']) in ID2K:
            a=age.setdefault(ID2K[int(r['LocID'])],[0,0,0]);s=int(r['AgeGrpStart']);a[0 if s<15 else 1 if s<65 else 2]+=float(r['PopTotal'])
for k,a in age.items():
    t=sum(a);D['pop'][k]['age']=[R(x/t*100) for x in a]

# ── UN WUP 2025 (국가별 도시 정의 기준) ──
def wup(f,sheet):
    ws=openpyxl.load_workbook(P('un','wup',f),read_only=True)[sheet];rows=list(ws.iter_rows(values_only=True))
    hi=next(i for i,r in enumerate(rows) if r and r[0]=='Index');hdr=[str(x) for x in rows[hi]]
    out={}
    for r in rows[hi+1:]:
        if isinstance(r[3],int) and r[3] in ID2K:out[ID2K[r[3]]]={int(hdr[j]):r[j] for j in range(10,len(hdr)) if hdr[j].isdigit()}
    return out
up=wup('WUP2025-F15-National-Definitions_percPop_by_category.xlsx','Urban')
upop=wup('WUP2025-F14-National-Definitions_Pop_by_category.xlsx','Urban');rpop=wup('WUP2025-F14-National-Definitions_Pop_by_category.xlsx','Rural')
ug=wup('WUP2025-F16-National-Definitions_Pop_rate_by_category.xlsx','Urban');rg=wup('WUP2025-F16-National-Definitions_Pop_rate_by_category.xlsx','Rural')
D['urban']={k:{'pct':{str(y):R(up[k][y]) for y in YRS},'urbanPop':R(upop[k][2025]/1000),'ruralPop':R(rpop[k][2025]/1000),
  'urbanGrowth':R(ug[k][2025],2),'ruralGrowth':R(rg[k][2025],2)} for k in ID2K.values()}

# ── UN DESA 국제 이주자 2024 ──
wb=openpyxl.load_workbook(P('un','undesa_pd_2024_ims_stock_by_sex_and_destination.xlsx'),read_only=True,data_only=True)
def desa(table,col):
    rows=list(wb[table].iter_rows(values_only=True));out={}
    for r in rows:
        loc=r[4] if table!='Table 2' else r[3]
        if isinstance(loc,int) and loc in ID2K and r[1]:out[ID2K[loc]]=r[col]
    return out
hdr=next(r for r in wb['Table 1'].iter_rows(values_only=True) if r and r[1] and str(r[1]).startswith('Region'));yc=list(hdr).index(2024)
mst=desa('Table 1',yc);mpct=desa('Table 3',yc)
wo=openpyxl.load_workbook(P('un','undesa_pd_2024_ims_stock_by_sex_destination_and_origin.xlsx'),read_only=True,data_only=True)['Table 1']
rows=list(wo.iter_rows(values_only=True));hh=next(r for r in rows if r and r[1] and str(r[1]).startswith('Region'));yo=list(hh).index(2024)
orig={k:{} for k in ID2K.values()}
for r in rows:
    if isinstance(r[4],int) and r[4] in ID2K and isinstance(r[6],int) and r[6] in ID2K:orig[ID2K[r[4]]][ID2K[r[6]]]=r[yo]
D['migrants']={k:{'stock':mst.get(k),'pct':R(mpct.get(k),2),'fromOrigin':orig[k]} for k in ID2K.values()}

# ── FAOSTAT 2023 지역 행 ──
FA={'af':['Africa'],'as':['Asia'],'eu':['Europe'],'an':['Northern America'],'oc':['Oceania'],'la':['Central America','South America','Caribbean'],'world':['World']}
PROD={'wheat':('15','Production','t',1e4,'밀'),'rice':('27','Production','t',1e4,'쌀(벼)'),'maize':('56','Production','t',1e4,'옥수수'),
      'cassava':('125','Production','t',1e4,'카사바'),'coffee':('656','Production','t',1e4,'커피 생두'),'cocoa':('661','Production','t',1e4,'카카오'),
      'tea':('667','Production','t',1e4,'차'),'rubber':('836','Production','t',1e4,'천연고무'),'sugarcane':('156','Production','t',1e4,'사탕수수'),
      'cotton':('328','Production','t',1e4,'씨앗 목화'),'olive':('260','Production','t',1e4,'올리브'),'grape':('560','Production','t',1e4,'포도'),
      'dates':('577','Production','t',1e4,'대추야자'),'beef':('867','Production','t',1e4,'소고기'),'pork':('1035','Production','t',1e4,'돼지고기'),
      'mutton':('977','Production','t',1e4,'양고기'),'cattle':('866','Stocks','An',1e4,'소 사육두수'),'pig':('1034','Stocks','An',1e4,'돼지 사육두수'),
      'sheep':('976','Stocks','An',1e4,'양 사육두수')}
def fao(zname,wanted,years):
    z=zipfile.ZipFile(P('fao',zname));nm=next(x for x in z.namelist() if x.endswith('(Normalized).csv'));out={}
    names={n for v in FA.values() for n in v}
    with io.TextIOWrapper(z.open(nm),encoding='latin-1') as h:
        for r in csv.DictReader(h):
            if r['Year'] in years and r['Area'] in names and (r['Item Code'],r['Element']) in wanted and r['Value']:
                out[(r['Item Code'],r['Element'],r['Area'],r['Year'])]=float(r['Value'])
    return out
want={(v[0],v[1]) for v in PROD.values()}
FY='2023'   # FAOSTAT 연도 — 수특과 같은 해
fp=fao('Production_Crops_Livestock_E_All_Data_(Normalized).zip',want,{FY})
def region(d,item,el,key,year,div):
    s=0;ok=False
    for a in FA[key]:
        v=d.get((item,el,a,year))
        if v is not None:s+=v;ok=True
    return R(s/div,1) if ok else None
D['food']={'prod':{}}
for key,(item,el,unit,div,label) in PROD.items():
    D['food']['prod'][key]={'label':label,'unit':'만 두' if unit=='An' else '만 톤','year':int(FY),
      'world':region(fp,item,el,'world',FY,div),**{k:region(fp,item,el,k,FY,div) for k in ID2K.values()}}
# 팜유
fp2=fao('Production_Crops_Livestock_E_All_Data_(Normalized).zip',{('257','Production')},{FY})
D['food']['prod']['palm']={'label':'팜유','unit':'만 톤','year':int(FY),'world':region(fp2,'257','Production','world',FY,1e4),**{k:region(fp2,'257','Production',k,FY,1e4) for k in ID2K.values()}}
tw=[('15','Export quantity'),('15','Import quantity'),('30','Export quantity'),('30','Import quantity'),('56','Export quantity'),('56','Import quantity')]
ft=fao('Trade_CropsLivestock_E_All_Data_(Normalized).zip',set(tw),{FY})
D['food']['trade']={}
for item,lab in(('15','밀'),('30','쌀(도정 환산)'),('56','옥수수')):
    D['food']['trade'][lab]={k:{'exp':region(ft,item,'Export quantity',k,FY,1e4),'imp':region(ft,item,'Import quantity',k,FY,1e4)} for k in ID2K.values()}
    D['food']['trade'][lab]['world']={'exp':region(ft,item,'Export quantity','world',FY,1e4),'imp':region(ft,item,'Import quantity','world',FY,1e4)}

# ── Ember 2024 전력 (나라 합산) ──
EM={'Coal':'석탄','Gas':'가스','Other Fossil':'기타 화석','Nuclear':'원자력','Hydro':'수력','Solar':'태양광','Wind':'풍력','Bioenergy':'바이오','Other Renewables':'기타 재생'}
el={k:{v:0.0 for v in EM} for k in ID2K.values()};tot={k:0.0 for k in ID2K.values()}
miss=set()
for r in csv.DictReader(open(P('other','ember_yearly_full_release_long_format.csv'),encoding='utf-8-sig')):
    if r['Year']!='2024' or r['Category']!='Electricity generation' or r['Unit']!='TWh' or r['Area type']!='Country or economy':continue
    k=iso2cont.get(r['ISO 3 code'])
    if not k:miss.add(r['Area']);continue
    if r['Subcategory']=='Fuel' and r['Variable'] in EM:el[k][r['Variable']]+=float(r['Value'] or 0)
D['power']={k:{'byFuel':{EM[v]:R(x,1) for v,x in el[k].items()},'total':R(sum(el[k].values()),1)} for k in el}
wt=sum(D['power'][k]['total'] for k in D['power'])
for k in D['power']:D['power'][k]['share']=R(D['power'][k]['total']/wt*100)
D['powerMissing']=sorted(miss)

# ── UNHCR 2024 (나라 합산) ──
ro={k:0 for k in ID2K.values()};rh={k:0 for k in ID2K.values()}
for r in json.load(open(P('unhcr','unhcr_2024.json'),encoding='utf-8')):
    n=int(r.get('refugees') or 0)
    a=iso2cont.get(r.get('coo_iso'));b=iso2cont.get(r.get('coa_iso'))
    if a:ro[a]+=n
    if b:rh[b]+=n
D['refugees']={'origin':ro,'host':rh,'year':2024}

# ── World Mining Data 2023 (나라 합산) ──
MIN={'Iron (Fe)':'철광석(철 함유량)','Bauxite':'보크사이트','Copper':'구리','Gold':'금','Silver':'은','Tin':'주석','Cobalt':'코발트','Manganese':'망간','Chromium (Cr2O3)':'크롬'}
WN={'Bosnia-Herzegovina':'BIH','Congo, Rep.':'COG','Congo, D.R.':'COD','Korea, North':'PRK','Korea, South':'KOR','Russia':'RUS',"Cote d'Ivoire":'CIV'}
wm=openpyxl.load_workbook(P('wmd','6.4_by_minerals.xlsx'),read_only=True,data_only=True);D['minerals']={}
for nm,lab in MIN.items():
    rows=list(wm[nm].iter_rows(values_only=True));hd=next(r for r in rows if r and r[0]=='Country');yc=list(hd).index('2023')
    s={k:0.0 for k in ID2K.values()};unit=rows[rows.index(hd)+1][1]
    for r in rows[rows.index(hd)+1:]:
        if not r[0] or r[0]=='Total' or r[yc] in(None,''):continue
        a3=WN.get(str(r[0]).strip())
        if a3 is None:
            try:a3=pycountry.countries.lookup(str(r[0]).strip()).alpha_3
            except LookupError:continue
        k=iso2cont.get(a3)
        if k:s[k]+=float(r[yc])
    t=sum(s.values())
    if unit=='kg':div,u=1000,'톤'
    elif t>=1e7:div,u=1e4,'만 톤'
    else:div,u=1,'톤'
    D['minerals'][lab]={'unit':u,'year':2023,'total':R(t/div,1),**{k:R(v/div,1) for k,v in s.items()},'share':{k:R(v/t*100) for k,v in s.items()} if t else {}}

# ── Energy Institute 2025판 권역 (EI 구분 그대로) ──
ei=openpyxl.load_workbook(P('ei','ei.xlsx'),read_only=True,data_only=True)
RG=[('Total North America','북아메리카(멕시코 포함)'),('Total S. & Cent. America','중남미'),('Total Europe','유럽'),('Total CIS','CIS'),
    ('Total Middle East','중동'),('Total Africa','아프리카'),('Total Asia Pacific','아시아태평양')]
def eitab(name):
    rows=list(ei[name].iter_rows(values_only=True))
    hi=next(i for i,r in enumerate(rows) if r and len(r)>1 and isinstance(r[1],(int,float)) and r[1]>=1960);yc=[j for j,x in enumerate(rows[hi]) if x==2024][0]
    return {str(r[0]).strip():r[yc] for r in rows[hi+1:] if r[0] and isinstance(r[yc],(int,float))}
EJ=277.78
D['energy']={'year':2024,'regions':['북아메리카(멕시코 포함)','중남미','유럽(러시아 포함)','건조 아시아(중동·CIS)','아프리카','아시아태평양'],'unit':'TWh','note':'Energy Institute Statistical Review 2025의 권역 구분(대륙과 경계가 다르다)'}
for key,sheet,mul in(('coalProd','Coal Production - EJ',EJ),('gasProd','Gas Production - EJ',EJ),('oilProd','Oil Production - tonnes',11.63),
                     ('coalCons','Coal Consumption - EJ',EJ),('gasCons','Gas Consumption - EJ',EJ),('oilCons','Oil Consumption - EJ',EJ)):
    t=eitab(sheet);w=t['Total World']
    row={lab:t[nm]*mul for nm,lab in RG if nm in t}
    # 러시아는 유럽으로 옮기고, 나머지 CIS와 중동은 '건조 아시아'로 묶는다(세계지리 교과의 지역 구분)
    ru=t.get('Russian Federation',0)*mul
    row['유럽(러시아 포함)']=row.pop('유럽',0)+ru
    row['건조 아시아(중동·CIS)']=row.pop('중동',0)+row.pop('CIS',0)-ru
    D['energy'][key]={'world':round(w*mul),**{k:round(v) for k,v in row.items()}}

# ── 경제 블록 — 회원국 값을 더한 합계. 회원 목록은 abyss/js/bloc-data.js 와 같은 것을 읽는다 ──
WD=os.path.join(os.path.dirname(os.path.abspath(OUT)),'world-data.js')
wtxt=open(WD,encoding='utf-8').read();WDATA=json.loads(re.search(r'^const WORLD_DATA=(\{.*\});\s*$',wtxt,flags=re.M).group(1))
btxt=open(os.path.join(os.path.dirname(os.path.abspath(OUT)),'bloc-data.js'),encoding='utf-8').read()
BL=[]
for m in re.finditer(r"\{id:'(\w+)',name:'([^']+)',short:'([^']+)',\s*members:\[([^\]]*)\],\s*intra:(\{[^}]*\}|null)",btxt):
    intra=None
    if m.group(5)!='null':
        mm=re.search(r'share:([\d.]+)',m.group(5));intra=float(mm.group(1)) if mm else None
    BL.append({'id':m.group(1),'name':m.group(2),'short':m.group(3),'members':re.findall(r"'(\w+)'",m.group(4)),'intra':intra})
# 인구(WPP 2025, 백만 명) · GDP(UN 통계국 2024, 억 달러)는 원천의 나라 값을 더한다
pop2={}
with gzip.open(P('un','WPP2024_Demographic_Indicators_Medium.csv.gz'),'rt',encoding='utf-8-sig') as h:
    for r in csv.DictReader(h):
        if r['Time']=='2025' and r['LocTypeName']=='Country/Area' and r['ISO3_code']:
            c=pycountry.countries.get(alpha_3=r['ISO3_code'])
            if c:pop2[c.alpha_2.lower()]=float(r['TPopulation1July'])/1000
gdp2={}
ws=openpyxl.load_workbook(P('unsd','Download-GDPcurrent-USD-countries.xlsx'),read_only=True,data_only=True).worksheets[0]
rr=list(ws.iter_rows(values_only=True));hi=next(i for i,r in enumerate(rr) if r and r[0]=='CountryID');yc2=list(rr[hi]).index(2024)
for r in rr[hi+1:]:
    if r[2]=='Gross Domestic Product (GDP)' and isinstance(r[0],int) and r[yc2]:
        c=pycountry.countries.get(numeric=str(r[0]).zfill(3))
        if c:gdp2[c.alpha_2.lower()]=float(r[yc2])/1e8
KEYS=['wh','ri','co','cas','cof','cca','tea','rub','pal','suc','ctn','oli','grp','dat','bef','prk','mtn','ct','pg','sh',
      'whx','whm','rix','rim','cox','com','cp','gp','op','cc','gc','oc','ox','om','cx','cm','gx','gm',
      'ior','bux','cop','gld','slv','tin','cbt','mng','chr','tx','tm','mva','rfo','rfh','mig']
D['blocs']=[]
for b in BL:
    agg={'pop':0.0,'gdp':0.0};cov={'pop':0,'gdp':0};el=[0.0]*10;elc=0
    for i in b['members']:
        if i in pop2:agg['pop']+=pop2[i];cov['pop']+=1
        if i in gdp2:agg['gdp']+=gdp2[i];cov['gdp']+=1
        w=WDATA.get(i,{})
        for k in KEYS:
            if k in w:agg[k]=agg.get(k,0)+w[k];cov[k]=cov.get(k,0)+1
        if w.get('el'):el=[a+b2 for a,b2 in zip(el,w['el'])];elc+=1
    agg={k:R(v,1) for k,v in agg.items()}
    D['blocs'].append({**b,'n':len(b['members']),'agg':agg,'cover':cov,'el':[R(x,1) for x in el],'elCover':elc})
# 지도 색칠용 — 대륙마다 속한 나라(ISO alpha-2 소문자, WPP 지역 계층 기준)
D['members']={}
for iso3,k in iso2cont.items():
    c=pycountry.countries.get(alpha_3=iso3)
    if c and k:D['members'].setdefault(k,[]).append(c.alpha_2.lower())
D['members']={k:sorted(v) for k,v in D['members'].items()}
D['blocNote']='회원국 값의 합계 — 인구 UN WPP 2025, GDP UN 통계국 2024, 나머지는 world-data.js(FAOSTAT 2023·EI 2024·WTO 2024 등)'

open(OUT,'w',encoding='utf-8').write('/* 대륙별 통계 — tools/world-update/continents.py 가 만든다(직접 고치지 말 것). 출처와 방법은 그 파일 머리말 */\nconst CONT_DATA='+json.dumps(D,ensure_ascii=False,separators=(',',':'))+';\n')
print('ok',os.path.getsize(OUT),'Ember 매칭 안 된 나라:',D['powerMissing'][:10])
for k,*_ in C:print(k,D['pop'][k]['series']['2025'],D['pop'][k]['age'],D['urban'][k]['pct']['2025'],D['migrants'][k]['stock'],D['power'][k]['total'],D['food']['prod']['wheat'][k])
